import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.chart import BarChart, Reference
import matplotlib.pyplot as plt
from io import BytesIO

def generate_excel_report(df, output_path):
    """Generate Excel report with charts"""
    # Create summary
    summary = df.groupby('Category')['Amount'].sum().reset_index()
    
    # Create workbook
    wb = Workbook()
    ws_trans = wb.active
    ws_trans.title = "Transactions"
    ws_summary = wb.create_sheet("Summary")
    
    # Write transaction data
    for r in dataframe_to_rows(df, index=False, header=True):
        ws_trans.append(r)
    
    # Write summary data
    for r in dataframe_to_rows(summary, index=False, header=True):
        ws_summary.append(r)
    
    # Create chart
    chart = BarChart()
    chart.title = "Spending by Category"
    chart.x_axis.title = "Category"
    chart.y_axis.title = "Amount"
    
    data = Reference(ws_summary, min_col=2, min_row=1, max_row=len(summary)+1)
    cats = Reference(ws_summary, min_col=1, min_row=2, max_row=len(summary)+1)
    
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws_summary.add_chart(chart, "D2")
    
    wb.save(output_path)